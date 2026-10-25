import os
import docx
from docx import Document
from docx.shared import Mm, Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, hex_color: str):
    """Установка фонового цвета ячейки таблицы."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    """Настройка внутренних отступов ячейки таблицы (в dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_page_number(run):
    """Вставка динамического поля номера страницы по ГОСТ 7.32."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)


def format_table_cell(
    cell,
    text: str,
    width=None,
    bg_color=None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    bold: bool = False,
    font_size: float = 10.0,
    font_name: str = 'Times New Roman',
    line_spacing: float = 1.15
):
    """
    Заполнение ячейки таблицы с гарантированным сбросом красной строки (first_line_indent = 0)
    и запретом выравнивания по ширине (JUSTIFY) для ВСЕХ абзацев ячейки.
    """
    cell.text = text
    if width is not None:
        cell.width = width
    if bg_color:
        set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    for p in cell.paragraphs:
        p.alignment = align
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = line_spacing
        for r in p.runs:
            r.font.name = font_name
            r.font.size = Pt(font_size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor(0, 0, 0)


def add_gost_paragraph(
    doc: Document,
    text: str = "",
    bold: bool = False,
    italic: bool = False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before: int = 0,
    space_after: int = 0,
    first_line_indent: float = 1.25,
    line_spacing: float = 1.5,
    font_size: int = 14
):
    """Добавление абзаца текста, оформленного строго по ГОСТ 7.32-2017."""
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(first_line_indent) if first_line_indent > 0 else Pt(0)

    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(font_size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)

    return p


def add_gost_heading(doc: Document, title: str, is_structural: bool = False):
    """
    Добавление заголовка по ГОСТ 7.32-2017:
    - Структурные заголовки (ВВЕДЕНИЕ, ЗАКЛЮЧЕНИЕ): по центру, прописными буквами, без точки.
    - Разделы: с абзацного отступа 1.25 см, полужирным, без точки.
    """
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5

    if is_structural:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(title.upper())
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(1.25)
        r = p.add_run(title)

    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_table_caption(doc: Document, number: int, title: str):
    """
    Оформление названия таблицы по ГОСТ 7.32-2017 (п. 6.8):
    Помещается над таблицей слева, без абзацного отступа, формат:
    Таблица N — Название таблицы
    """
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    r = p.add_run(f"Таблица {number} — {title}")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = False
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def create_report():
    doc = Document()

    # 1. ПАРАМЕТРЫ СТРАНИЦЫ ПО ГОСТ 7.32-2017
    # Формат А4, левое поле 30 мм, правое 15 мм, верхнее 20 мм, нижнее 20 мм
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)

    # Титульный лист без номера, нумерация со 2-й страницы внизу по центру
    section.different_first_page_header_footer = True
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = 0
    footer_p.paragraph_format.space_after = 0
    footer_p.paragraph_format.first_line_indent = Pt(0)
    footer_run = footer_p.add_run()
    footer_run.font.name = 'Times New Roman'
    footer_run.font.size = Pt(11)
    footer_run.font.color.rgb = RGBColor(0, 0, 0)
    add_page_number(footer_run)

    # Базовый стиль шрифта (без принудительного отступа, чтобы не ломать таблицы)
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(0)
    normal_style.paragraph_format.first_line_indent = Pt(0)
    normal_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # =========================================================================
    # ТИТУЛЬНЫЙ ЛИСТ (ГОСТ 7.32-2017)
    # =========================================================================
    add_gost_paragraph(
        doc,
        "МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ\n"
        "Федеральное государственное автономное образовательное учреждение высшего образования\n"
        "«Новосибирский национальный исследовательский государственный университет»\n"
        "Высший колледж информатики",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=0,
        line_spacing=1.1,
        font_size=11,
        space_after=48
    )

    add_gost_paragraph(
        doc,
        "ОТЧЕТ\n"
        "ПО ЛАБОРАТОРНОЙ РАБОТЕ № 2\n"
        "по дисциплине «Проектирование и тестирование программных модулей»\n"
        "Тема: «Юнит-тестирование проектов и автоматизация проверок»",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=0,
        line_spacing=1.3,
        font_size=14,
        space_before=24,
        space_after=60
    )

    # Блок подписей студента и преподавателя (справа)
    add_gost_paragraph(
        doc,
        "Выполнил:\n"
        "студент Романов Р. О.\n"
        "Специальность: 09.02.07 Информационные\n"
        "системы и программирование\n\n"
        "Проверил:\n"
        "преподаватель Кошман В. Е.\n"
        "_________________________\n",
        bold=False,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_line_indent=0,
        line_spacing=1.15,
        font_size=12,
        space_before=20,
        space_after=140
    )

    # Город и год - строго снизу по центру без красной строки
    p_city = doc.add_paragraph()
    p_city.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_city.paragraph_format.first_line_indent = Pt(0)
    p_city.paragraph_format.space_before = Pt(40)
    p_city.paragraph_format.space_after = Pt(0)
    p_city.paragraph_format.line_spacing = 1.0
    r_city = p_city.add_run("Новосибирск, 2026")
    r_city.font.name = 'Times New Roman'
    r_city.font.size = Pt(12)
    r_city.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()

    # =========================================================================
    # ВВЕДЕНИЕ
    # =========================================================================
    add_gost_heading(doc, "ВВЕДЕНИЕ", is_structural=True)

    add_gost_paragraph(
        doc,
        "Модульное тестирование (Unit Testing) представляет собой фундаментальный метод обеспечения качества "
        "программного обеспечения, позволяющий проверять корректность работы отдельных функций, методов и классов "
        "в изолированной среде. Регулярный запуск автоматизированных тестов снижает риски возникновения регрессионных дефектов "
        "и повышает предсказуемость поведения программного комплекса."
    )

    add_gost_paragraph(
        doc,
        "Цель работы — освоение методологии разработки юнит-тестов на базе стандартной библиотеки unittest в языке Python, "
        "применение техник тест-дизайна (эквивалентное разбиение, анализ граничных значений), а также обнаружение, "
        "локализация и протоколирование дефектов логики в предоставленном учебном программном модуле."
    )

    add_gost_paragraph(doc, "Для достижения поставленной цели в рамках работы решены следующие задачи:")

    tasks = [
        "разработать комплект независимых юнит-тестов (не менее 10) для собственного программного метода валидации учетных данных из лабораторной работы № 1;",
        "проанализировать исходный код и бизнес-требования к предоставленному модулю расчета стоимости и сроков транспортировки отправлений (src/delivery_service.py);",
        "спроектировать комплекс юнит-тестов (не менее 10) для проверки всех классов эквивалентности, граничных значений и исключительных ситуаций в модуле доставки;",
        "выполнить прогон разработанных наборов тестов с применением механизма автоматического обнаружения (discover) фреймворка unittest;",
        "зафиксировать упавшие тестовые сценарии, локализовать программные ошибки в коде и сформулировать обоснованные варианты их исправления."
    ]
    for idx, task in enumerate(tasks):
        terminator = "." if idx == len(tasks) - 1 else ";"
        add_gost_paragraph(
            doc,
            f"- {task}{terminator}",
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line_indent=1.25,
            line_spacing=1.5
        )

    # =========================================================================
    # 1 ПУНКТ А: СТАТИСТИКА ПРОХОЖДЕНИЯ ТЕСТОВ
    # =========================================================================
    add_gost_heading(doc, "1 Общее количество тестов и итоговая статистика прохождения (Пункт А)")

    add_gost_paragraph(
        doc,
        "В соответствии с требованиями к лабораторной работе № 2 были подготовлены два автономных тестовых комплекта, "
        "размещенных в каталоге tests/. Комплекты покрывают как собственный метод регистрации пользователей из лабораторной "
        "работы № 1, так и предоставленный модуль расчета параметров доставки. Всего разработан 45 юнит-тестов."
    )

    add_table_caption(doc, 1, "Итоговая статистика прохождения юнит-тестов проекта")

    # Таблица 1 по ГОСТу
    table_stat = doc.add_table(rows=4, cols=6)
    table_stat.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_stat.style = 'Table Grid'

    col_widths_t1 = [Inches(1.8), Inches(1.8), Inches(0.8), Inches(0.9), Inches(0.9), Inches(0.9)]

    t1_headers = [
        "Тестируемый модуль",
        "Файл тестов",
        "Всего тестов",
        "Пройдено (OK)",
        "Упало (Fail/Err)",
        "Успешность"
    ]

    for col_idx, h in enumerate(t1_headers):
        cell = table_stat.cell(0, col_idx)
        format_table_cell(
            cell,
            text=h,
            width=col_widths_t1[col_idx],
            bg_color="F2F2F2",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            font_size=10.5
        )

    t1_data = [
        ("ЛР1: Регистрация пользователей", "tests/test_my_project.py", "22", "22", "0", "100.0%"),
        ("ЛР2: Расчет стоимости доставки", "tests/test_delivery.py", "23", "19", "4", "82.6%"),
        ("ИТОГО ПО ПРОЕКТУ", "Все модули каталога tests/", "45", "41", "4", "91.1%")
    ]

    for row_idx, row_data in enumerate(t1_data, start=1):
        is_total = (row_idx == 3)
        bg = "EBEBEB" if is_total else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = table_stat.cell(row_idx, col_idx)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx >= 2 else WD_ALIGN_PARAGRAPH.LEFT
            format_table_cell(
                cell,
                text=val,
                width=col_widths_t1[col_idx],
                bg_color=bg,
                align=align,
                bold=is_total,
                font_size=10.0
            )

    add_gost_paragraph(
        doc,
        "Набор тестов собственного проекта (tests/test_my_project.py) содержит 22 теста и проверяет соответствие "
        "входных параметров маскам телефонных номеров (+x-xxx-xxx-xxxx), формату адресов электронной почты по стандарту RFC, "
        "длине и допустимым символам строковых идентификаторов, фильтрацию по черному списку служебных учетных записей, "
        "а также обязательное наличие в пароле прописных и строчных букв кириллицы, цифр и спецсимволов. "
        "Все 22 теста успешно завершились (статус OK), подтверждая корректность ранее реализованной бизнес-логики.",
        space_before=8
    )

    add_gost_paragraph(
        doc,
        "Набор тестов модуля доставки (tests/test_delivery.py) включает 23 тестовых метода, охватывающих базовый тариф (200 руб.), "
        "километражную составляющую (5 руб./км), весовые коэффициенты (1.2 при весе от 5 до 20 кг; 1.5 при весе от 20 до 50 кг), "
        "дополнительные наценки по категории груза («хрупкий» +300 руб., «опасный» +1000 руб.), экспресс-доставку, расчет даты "
        "транспортировки от фиксированной даты (2026-09-03) и граничные физические ограничения (0.1 кг, 50.0 кг, 1 км, 5000 км)."
    )

    # =========================================================================
    # 2 ПУНКТ Б: КОЛИЧЕСТВО УПАВШИХ ТЕСТОВ ДОСТАВКИ
    # =========================================================================
    add_gost_heading(doc, "2 Точное количество упавших тестов для модуля доставки (Пункт Б)")

    add_gost_paragraph(
        doc,
        "При автоматизированном прогоне набора юнит-тестов tests/test_delivery.py на предоставленной исходной реализации "
        "модуля src/delivery_service.py зафиксированы следующие результаты:"
    )

    add_gost_paragraph(doc, "- точное количество не прошедших (упавших) тестов: 4 (четыре) теста из 23;")
    add_gost_paragraph(doc, "- количество падений с ошибкой валидации утверждения (AssertionError / FAIL): 3 теста;")
    add_gost_paragraph(doc, "- количество падений с необработанным исключением исполнения (TypeError / ERROR): 1 тест.")

    add_table_caption(doc, 2, "Сводный перечень упавших тестов модуля доставки src/delivery_service.py")

    # Таблица 2 по ГОСТу
    table_fails = doc.add_table(rows=5, cols=4)
    table_fails.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_fails.style = 'Table Grid'

    col_widths_t2 = [Inches(0.6), Inches(2.3), Inches(1.2), Inches(2.9)]
    t2_headers = ["№", "Название теста", "Тип ошибки", "Текст ошибки и характер сбоя"]

    for col_idx, h in enumerate(t2_headers):
        cell = table_fails.cell(0, col_idx)
        format_table_cell(
            cell,
            text=h,
            width=col_widths_t2[col_idx],
            bg_color="F2F2F2",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            font_size=10.5
        )

    t2_data = [
        (
            "1",
            "test_express_delivery_cost_surcharge",
            "AssertionError (FAIL)",
            "350 not greater than 700. Стоимость экспресс-доставки оказалась в 2 раза ниже базовой (скидка вместо наценки за срочность)."
        ),
        (
            "2",
            "test_express_delivery_minimum_one_day",
            "AssertionError (FAIL)",
            "'2026-09-03' != '2026-09-04'. Срок доставки составил 0 дней (день в день), что нарушает требование о минимальном сроке транспортировки."
        ),
        (
            "3",
            "test_package_type_case_insensitivity",
            "AssertionError (FAIL)",
            "-1 != 700. Значение 'Обычный' с заглавной буквы привело к ложному возврату ошибки из-за регистрозависимой проверки списка допустимых типов."
        ),
        (
            "4",
            "test_invalid_input_types_graceful_handling",
            "TypeError (ERROR)",
            "TypeError: '<' not supported between instances of 'str' and 'float'. Аварийный сбой программы при передаче строкового веса 'десять'."
        )
    ]

    for row_idx, data in enumerate(t2_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table_fails.cell(row_idx, col_idx)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            bold = (col_idx == 1)
            format_table_cell(
                cell,
                text=text,
                width=col_widths_t2[col_idx],
                bg_color="FFFFFF",
                align=align,
                bold=bold,
                font_size=9.5
            )

    # =========================================================================
    # 3 ПУНКТ В: ЛОКАЛИЗАЦИЯ ДЕФЕКТОВ
    # =========================================================================
    add_gost_heading(doc, "3 Локализация аномалий и дефектов в коде модуля доставки (Пункт В)")

    add_gost_paragraph(
        doc,
        "По результатам анализа сообщений об ошибках и трассировок стека вызовов фреймворка unittest была проведена "
        "детальная локализация дефектов в исходном коде src/delivery_service.py. Для каждой обнаруженной аномалии определен "
        "точный номер некорректно работающей строки, сформулирована причина дефекта в логике разработчика и предложен "
        "исправленный вариант программного кода. Результаты локализации приведены в таблице 3."
    )

    add_table_caption(doc, 3, "Локализация дефектов в модуле src/delivery_service.py и варианты их исправления")

    # Таблица 3 по ГОСТу
    table_bugs = doc.add_table(rows=5, cols=4)
    table_bugs.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_bugs.style = 'Table Grid'

    col_widths_t3 = [Inches(1.8), Inches(0.8), Inches(2.3), Inches(2.1)]
    t3_headers = ["Название теста", "Строка", "Описание дефекта в логике", "Вариант исправления кода"]

    for col_idx, h in enumerate(t3_headers):
        cell = table_bugs.cell(0, col_idx)
        format_table_cell(
            cell,
            text=h,
            width=col_widths_t3[col_idx],
            bg_color="F2F2F2",
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bold=True,
            font_size=10.5
        )

    t3_data = [
        (
            "test_express_delivery_cost_surcharge",
            "35",
            "Искажение расчета экспресс-тарифа: разработчик применил умножение на 0.5 (total_cost *= 0.5), что дает скидку 50% вместо положенной наценки за срочность.",
            "if is_express:\n    total_cost *= 1.5\n# Наценка 50% к стоимости"
        ),
        (
            "test_express_delivery_minimum_one_day",
            "43",
            "Логическая ошибка целочисленного деления: при distance < 1000 км базовый срок равен 1 дню. Операция 1 // 2 дает 0 дней, приводя к фиктивной доставке в день отправления.",
            "if is_express:\n    days_needed = max(1, days_needed // 2)\n# Не менее 1 дня"
        ),
        (
            "test_package_type_case_insensitivity",
            "15",
            "Чувствительность к регистру символов: проверка 'package_type not in valid_types' отклоняет валидные типы 'Обычный', 'Хрупкий' с заглавной буквы.",
            "clean_type = str(package_type).strip().lower()\nif clean_type not in valid_types:\n    return -1, '0000-00-00'"
        ),
        (
            "test_invalid_input_types_graceful_handling",
            "11",
            "Отсутствие строгой типизации: при передаче нечисловых параметров (например, weight='десять') функция аварийно падает с TypeError при сравнении weight < 0.1.",
            "if not isinstance(weight, (int, float)) or not isinstance(distance, int) or isinstance(distance, bool):\n    return -1, '0000-00-00'"
        )
    ]

    for row_idx, data in enumerate(t3_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table_bugs.cell(row_idx, col_idx)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            font_name = 'Consolas' if col_idx == 3 else 'Times New Roman'
            font_size = 8.5 if col_idx == 3 else 9.5
            bold = (col_idx == 0)
            format_table_cell(
                cell,
                text=text,
                width=col_widths_t3[col_idx],
                bg_color="FFFFFF",
                align=align,
                bold=bold,
                font_size=font_size,
                font_name=font_name
            )

    # =========================================================================
    # 4 МЕТОДОЛОГИЯ ТЕСТ-ДИЗАЙНА
    # =========================================================================
    add_gost_heading(doc, "4 Методология тест-дизайна и анализ тестового покрытия")

    add_gost_paragraph(
        doc,
        "Для формирования полного и неизбыточного набора тестов были применены классические техники тест-дизайна, "
        "рекомендованные международным стандартом ISTQB и отечественными методиками тестирования ПО:"
    )

    add_gost_paragraph(
        doc,
        "1. Эквивалентное разбиение (Equivalence Partitioning). Область допустимых значений параметров разбита на непересекающиеся классы:\n"
        "- по весу посылки: класс без наценки (0.1–5.0 кг), класс умеренной надбавки 20% (5.1–19.9 кг), класс повышенной надбавки 50% (20.0–50.0 кг);\n"
        "- по категории груза: «обычный» (0 руб.), «хрупкий» (+300 руб.), «опасный» (+1000 руб.), недопустимый тип (возврат ошибки -1);\n"
        "- по срочности транспортировки: стандартная доставка (is_express=False) и срочная доставка (is_express=True)."
    )

    add_gost_paragraph(
        doc,
        "2. Анализ граничных значений (Boundary Value Analysis). Проверены точки переходов и предельные ограничения физической модели:\n"
        "- нижняя граница веса: 0.1 кг (валидный минимум), 0.05 кг (отказ по нижней границе), -5.0 кг (отрицательный вес);\n"
        "- верхняя граница веса: 50.0 кг (валидный максимум), 50.1 кг (отказ по верхней границе);\n"
        "- пороги коэффициентов тарификации: ровно 5.0 кг (граница базового тарифа), ровно 20.0 кг (граница тарифа 1.5);\n"
        "- ограничения расстояния: 1 км (минимальная дистанция), 0 км и -100 км (недопустимые значения), 5000 км (максимальная дистанция), 5001 км (превышение лимита)."
    )

    add_gost_paragraph(
        doc,
        "3. Негативное тестирование и проверка отказоустойчивости (Negative & Robustness Testing). Проверена реакция "
        "программных модулей на передачу непредусмотренных типов данных (строки вместо чисел, булевы значения, пустые структуры), "
        "что позволило выявить дефект отсутствия предварительной верификации типов входных аргументов."
    )

    # =========================================================================
    # ЗАКЛЮЧЕНИЕ
    # =========================================================================
    add_gost_heading(doc, "ЗАКЛЮЧЕНИЕ", is_structural=True)

    add_gost_paragraph(
        doc,
        "В ходе выполнения лабораторной работы № 2 была всесторонне исследована методология юнит-тестирования программных "
        "модулей на языке Python с применением стандартного фреймворка unittest. Разработанный тестовый комплекс из 45 тестов "
        "позволил подтвердить 100% стабильность собственного модуля регистрации пользователей (ЛР1) и локализовать 4 критических "
        "дефекта в предоставленном модуле расчета доставки (src/delivery_service.py)."
    )

    add_gost_paragraph(
        doc,
        "Опыт локализации ошибок продемонстрировал, что применение методов эквивалентного разбиения и анализа граничных значений "
        "позволяет находить как смысловые ошибки бизнес-логики (скидка вместо наценки за срочность), так и граничные аномалии "
        "целочисленного деления и исключения времени выполнения, обеспечивая высокий уровень надежности программного продукта."
    )

    # =========================================================================
    # СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ
    # =========================================================================
    add_gost_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", is_structural=True)

    sources = [
        "ГОСТ 7.32-2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления. — М.: Стандартинформ, 2018. — 32 с.",
        "ГОСТ Р 7.0.5-2008. Система стандартов по информации, библиотечному и издательскому делу. Библиографическая ссылка. Общие требования и правила составления. — М.: Стандартинформ, 2008. — 20 с.",
        "Документация модуля unittest в Python 3 [Электронный ресурс]. — Режим доступа: https://docs.python.org/3/library/unittest.html (дата обращения: 22.10.2026).",
        "Куликов С. Тестирование программного обеспечения. Базовый курс. — Минск: Четыре четверти, 2020. — 312 с.",
        "Бейзер Б. Тестирование черного ящика. Технологии функционального тестирования программного обеспечения и систем. — СПб.: Питер, 2004. — 320 с."
    ]

    for idx, s in enumerate(sources, start=1):
        add_gost_paragraph(
            doc,
            f"{idx}. {s}",
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line_indent=1.25,
            line_spacing=1.5,
            font_size=12
        )

    output_path = "Отчет_ЛР2_Романов_Р_О.docx"
    doc.save(output_path)
    print(f"Отчет по ГОСТ 7.32 успешно создан: {output_path}")


if __name__ == "__main__":
    create_report()
